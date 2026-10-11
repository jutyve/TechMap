import io
import os
import re
import zipfile
from typing import Optional
from xml.etree import ElementTree as ET

from docx import Document


def normalize_extracted_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_docx_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))
    lines = []

    def add(text: Optional[str]):
        text = (text or "").strip()
        if text:
            lines.append(text)

    # Body paragraphs
    for paragraph in document.paragraphs:
        add(paragraph.text)

    # Tables in the body
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                add(cell.text)

    # Headers / footers for each section
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            add(paragraph.text)
        for paragraph in section.footer.paragraphs:
            add(paragraph.text)

    return "\n".join(lines)


def extract_pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data), strict=False)
    pages = []

    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")

    return "\n".join(pages)


def extract_odt_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml_data = archive.read("content.xml")

    root = ET.fromstring(xml_data)
    text_ns = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
    paragraphs = []

    for element in root.iter():
        if element.tag not in (
            f"{{{text_ns}}}p",
            f"{{{text_ns}}}h",
        ):
            continue

        pieces = []

        for node in element.iter():
            if node.text:
                pieces.append(node.text)

            if node.tag == f"{{{text_ns}}}tab":
                pieces.append("\t")
            elif node.tag == f"{{{text_ns}}}line-break":
                pieces.append("\n")

            if node.tail:
                pieces.append(node.tail)

        line = "".join(pieces).strip()

        if line:
            paragraphs.append(line)

    return "\n".join(paragraphs)


def extract_text(filename: str, data: bytes) -> str:
    ext = os.path.splitext(filename.lower())[1]

    try:
        if ext in (".txt", ".md"):
            text = data.decode("utf-8-sig", errors="replace")
        elif ext == ".pdf":
            text = extract_pdf_text(data)
        elif ext == ".docx":
            text = extract_docx_text(data)
        elif ext == ".odt":
            text = extract_odt_text(data)
        else:
            return ""

        return normalize_extracted_text(text)

    except Exception as exc:
        print(
            f"Text extraction failed for {filename}: "
            f"{type(exc).__name__}: {exc}"
        )
        return ""
